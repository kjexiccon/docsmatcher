
import streamlit as st
import pandas as pd
import fitz
import docx
import tempfile
from fpdf import FPDF

st.set_page_config(page_title="Document Verifier – v6", layout="wide")
st.title("📦 SHEESH SPICES Document Verifier – v6 (Multi-Doc + Mismatch Codes)")

st.markdown("Upload your master file (Invoice or Packing List), then upload one or more documents to compare. All mismatches are flagged with reason codes.")

master_file = st.file_uploader("📘 Upload Master Excel", type=["xlsx"])
compare_files = st.file_uploader("📄 Upload Other Documents (COO, PHYTO, FUMI, B/L)", type=["pdf", "docx"], accept_multiple_files=True)

def extract_text(file):
    if file.name.endswith(".pdf"):
        text = ""
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(file.read())
            tmp.flush()
            doc = fitz.open(tmp.name)
            for page in doc:
                text += page.get_text()
        return text.lower()
    else:
        doc = docx.Document(file)
        return "\n".join([p.text for p in doc.paragraphs]).lower()

def compare_structured_rows(df, doc_text):
    results = []
    for idx, row in df.iterrows():
        product = row.get("Product", "").strip().lower()
        hs = row.get("HS Code", "").strip().lower()
        qty = str(row.get("Quantity", "")).strip().lower()

        row_result = {"Row": idx+2, "Product": product}
        issues = []

        if product not in doc_text:
            issues.append("PNM")  # Product Name Mismatch
        if hs and hs not in doc_text:
            issues.append("HSC")
        if qty and qty not in doc_text:
            issues.append("QTY")

        if not issues:
            row_result["Status"] = "✔ Matched"
        else:
            row_result["Status"] = "❌ " + ", ".join(issues)

        results.append(row_result)
    return pd.DataFrame(results)

def generate_pdf_report(df, title="Mismatch Report"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt=title, ln=True, align='C')
    pdf.ln(10)

    for idx, row in df.iterrows():
        line = f"Row {row['Row']} | Product: {row['Product']} | {row['Status']}"
        pdf.multi_cell(0, 10, txt=line)
    tmp_path = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf").name
    pdf.output(tmp_path)
    return tmp_path

if master_file and compare_files:
    df_master = pd.read_excel(master_file).fillna("").astype(str)
    st.markdown("### 🧾 Parsed Master Data")
    st.dataframe(df_master.head(), use_container_width=True)

    all_results = []
    for doc in compare_files:
        st.markdown(f"## 📄 Comparing with: `{doc.name}`")
        doc_text = extract_text(doc)
        result_df = compare_structured_rows(df_master, doc_text)

        show_only_mismatches = st.checkbox(f"🔍 Show only mismatches for `{doc.name}`", key=doc.name)
        display_df = result_df[result_df["Status"].str.startswith("❌")] if show_only_mismatches else result_df
        st.dataframe(display_df, use_container_width=True)

        all_results.append((doc.name, result_df))

        # PDF Download
        if not result_df[result_df["Status"].str.startswith("❌")].empty:
            pdf_path = generate_pdf_report(result_df[result_df["Status"].str.startswith("❌")], title=f"Mismatch Report - {doc.name}")
            with open(pdf_path, "rb") as f:
                st.download_button(f"📥 Download PDF Report for `{doc.name}`", data=f, file_name=f"report_{doc.name}.pdf", mime="application/pdf")

else:
    st.info("Upload master Excel and at least one document to start comparison.")
