
import streamlit as st
import pandas as pd
import fitz
import docx
import tempfile
from fpdf import FPDF

st.set_page_config(page_title="SHEESH SPICES Verifier – v9", layout="wide")
st.title("📦 SHEESH SPICES Document Verifier – v9")

st.markdown("Upload your master invoice/packing list and compare across COO, PHYTO, FUMI, and B/L.")
st.sidebar.header("📂 Upload Files")

master_file = st.sidebar.file_uploader("Primary Invoice or Packing List (Excel)", type=["xlsx"])
compare_files = st.sidebar.file_uploader("Compare Docs (COO, Phyto, Fumi, B/L)", type=["pdf", "docx"], accept_multiple_files=True)

def extract_text(file):
    if file.name.endswith(".pdf"):
        text = ""
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(file.read())
            tmp.flush()
            doc = fitz.open(tmp.name)
            for page in doc:
                text += page.get_text()
        return text.upper()
    else:
        doc = docx.Document(file)
        return "\n".join([p.text for p in doc.paragraphs]).upper()

def find_columns(df):
    header_map = {}
    for col in df.columns:
        col_lower = str(col).lower()
        if "product" in col_lower or "item" in col_lower:
            header_map["Product"] = col
        elif "hs" in col_lower:
            header_map["HS Code"] = col
        elif "qty" in col_lower:
            header_map["Quantity"] = col
    return header_map

def parse_master_rows(df):
    header_map = find_columns(df)
    rows = []
    for _, row in df.iterrows():
        prod = row.get(header_map.get("Product", ""), "").strip().upper()
        hs = row.get(header_map.get("HS Code", ""), "").strip()
        qty = str(row.get(header_map.get("Quantity", ""), "")).strip()
        if prod:
            rows.append({"Product": prod, "HS Code": hs, "Quantity": qty})
    return rows

def compare_field(prod_val, doc_text):
    return prod_val and prod_val in doc_text

def summarize_product(product, doc_texts):
    result = {"Product": product["Product"], "HS Code": product["HS Code"], "Quantity": product["Quantity"]}
    mismatches = []
    for name, text in doc_texts.items():
        status = []
        if not compare_field(product["Product"], text):
            status.append("PNM")
        if not compare_field(product["HS Code"], text):
            status.append("HSC")
        if not compare_field(product["Quantity"], text):
            status.append("QTY")
        result[name] = ", ".join(status) if status else "✅"
        mismatches.extend(status)
    result["Status"] = "✅" if not mismatches else "❌ " + ", ".join(set(mismatches))
    return result

def export_pdf(results):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, "SHEESH SPICES Document Comparison Report", ln=True, align="C")
    pdf.ln(10)
    for r in results:
        pdf.cell(200, 10, f"{r['Product']} | Status: {r['Status']}", ln=True)
    out_path = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf").name
    pdf.output(out_path)
    return out_path

if master_file and compare_files:
    st.success("✅ Files uploaded.")
    df = pd.read_excel(master_file).fillna("").astype(str)
    products = parse_master_rows(df)
    doc_texts = {doc.name: extract_text(doc) for doc in compare_files}
    results = [summarize_product(prod, doc_texts) for prod in products]
    df_result = pd.DataFrame(results)

    st.markdown("## 📊 Summary Dashboard")
    total = len(df_result)
    mismatches = df_result[df_result["Status"].str.startswith("❌")]
    st.write(f"**Total Products:** {total}")
    st.write(f"**Mismatches:** {len(mismatches)}")

    st.markdown("## 🔍 Detailed Comparison")
    st.dataframe(df_result, use_container_width=True)

    csv_data = df_result.to_csv(index=False).encode()
    st.download_button("📥 Download CSV Report", csv_data, "comparison_report.csv", "text/csv")

    pdf_path = export_pdf(results)
    with open(pdf_path, "rb") as f:
        st.download_button("🖨️ Download PDF Summary", f, file_name="comparison_report.pdf", mime="application/pdf")
else:
    st.info("Upload master file and comparison documents to begin.")
